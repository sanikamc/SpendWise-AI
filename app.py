"""
app.py - SpendWise AI Main Application
======================================
Smart Expense Anomaly Detection Web Application
Domain 5 - AI for Finance and Governance (College CSE Mini-Project)

Author: Senior Python / ML / Full-Stack Engineer
"""

import csv
from functools import wraps
import io
import json
import os
import secrets
from calendar import monthrange
from datetime import date, datetime, timedelta
from typing import Any, Dict, List, Optional, Tuple

from flask import (
    Flask,
    abort,
    flash,
    jsonify,
    redirect,
    render_template,
    request,
    Response,
    session,
    url_for,
)
from sqlalchemy import extract, func

from config import Config
from database import Budget, Expense, User, db
from model import detector


def format_csv_date(dt_val: Any) -> str:
    """
    Formats a transaction date into a consistent, human-readable YYYY-MM-DD string.
    Ensures dates are formatted date strings rather than Excel serial numbers
    (e.g. 46926) or ambiguous datetime representations.
    """
    if dt_val is None:
        return ""
    if isinstance(dt_val, (date, datetime)):
        return dt_val.strftime("%Y-%m-%d")
    s = str(dt_val).strip()
    if not s:
        return ""
    # Check if value is a numeric Excel date serial number (e.g. 46926 or "46926")
    try:
        num = float(s)
        if 1 <= num <= 200000:
            converted = date(1899, 12, 30) + timedelta(days=int(num))
            return converted.strftime("%Y-%m-%d")
    except (ValueError, TypeError, OverflowError):
        pass
    # Try parsing ISO or standard date formats
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except (ValueError, TypeError):
            continue
    return s[:10] if len(s) >= 10 else s


def create_app(config_class=Config) -> Flask:
    """Application factory for SpendWise AI."""
    app = Flask(__name__)
    app.config.from_object(config_class)

    # Ensure instance folder exists
    os.makedirs(app.config["INSTANCE_PATH"], exist_ok=True)

    # Initialize extensions
    db.init_app(app)

    with app.app_context():
        # Create database tables if they do not exist
        db.create_all()

    # -------------------------------------------------------------
    # Authentication & Session Helpers
    # -------------------------------------------------------------
    def get_current_user() -> Optional[User]:
        """Returns the currently authenticated User instance or None."""
        user_id = session.get("user_id")
        if user_id:
            return db.session.get(User, user_id)
        return None

    def login_required(f):
        """Decorator to restrict access to authenticated users."""
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if "user_id" not in session or not get_current_user():
                flash("Please sign in to access your personal dashboard.", "warning")
                return redirect(url_for("login", next=request.path))
            return f(*args, **kwargs)
        return decorated_function

    def generate_csrf_token() -> str:
        """Generates or retrieves a secure CSRF token for the session."""
        if "_csrf_token" not in session:
            session["_csrf_token"] = secrets.token_hex(16)
        return session["_csrf_token"]

    @app.before_request
    def csrf_protect():
        """Validates CSRF tokens on all state-changing HTTP requests."""
        if request.method in ["POST", "PUT", "DELETE", "PATCH"]:
            # Bypass CSRF in testing environment if configured
            if not app.config.get("WTF_CSRF_ENABLED", True):
                return
            token = request.form.get("csrf_token") or request.headers.get("X-CSRFToken")
            expected = session.get("_csrf_token")
            if not expected or not token or not secrets.compare_digest(str(token), str(expected)):
                abort(400, description="Invalid or missing CSRF security token.")

    # Register Context Processors for Templates
    @app.context_processor
    def inject_global_data():
        current_u = get_current_user()
        if current_u:
            demo_count = Expense.query.filter_by(user_id=current_u.id, is_demo=True).count()
            total_count = Expense.query.filter_by(user_id=current_u.id).count()
            pending_anomalies = Expense.query.filter_by(
                user_id=current_u.id, is_anomaly=True, review_status="pending"
            ).count()
        else:
            demo_count = 0
            total_count = 0
            pending_anomalies = 0

        return {
            "app_name": app.config["APP_NAME"],
            "app_subtitle": app.config["APP_SUBTITLE"],
            "currency_symbol": app.config["CURRENCY_SYMBOL"],
            "categories": app.config["EXPENSE_CATEGORIES"],
            "payment_methods": app.config["PAYMENT_METHODS"],
            "current_user": current_u,
            "csrf_token": generate_csrf_token,
            "has_demo_data": demo_count > 0,
            "demo_count": demo_count,
            "total_expense_count": total_count,
            "pending_anomaly_count": pending_anomalies,
            "current_year": date.today().year,
            "current_month": date.today().month,
        }

    # -------------------------------------------------------------
    # Error Handlers
    # -------------------------------------------------------------
    @app.errorhandler(400)
    def bad_request_error(error):
        return render_template("errors/400.html", message=str(error)), 400

    @app.errorhandler(403)
    def forbidden_error(error):
        return render_template("errors/403.html"), 403

    @app.errorhandler(404)
    def not_found_error(error):
        return render_template("errors/404.html"), 404

    @app.errorhandler(500)
    def internal_error(error):
        db.session.rollback()
        return render_template("errors/500.html"), 500

    # -------------------------------------------------------------
    # Helper Utilities
    # -------------------------------------------------------------
    def get_selected_period(req) -> Tuple[int, int]:
        """Extracts (year, month) from query params or defaults to current date."""
        today = date.today()
        try:
            year = int(req.args.get("year", today.year))
            month = int(req.args.get("month", today.month))
            if not (1 <= month <= 12 and 2000 <= year <= 2100):
                year, month = today.year, today.month
        except (ValueError, TypeError):
            year, month = today.year, today.month
        return year, month

    def retrain_model_on_database(user_id: int):
        """Refreshes IsolationForest on the specified user's database records."""
        user_expenses = (
            Expense.query.filter_by(user_id=user_id)
            .order_by(Expense.date.asc())
            .all()
        )
        if not user_expenses:
            return {"status": "empty", "message": "No transactions in database."}

        records = [e.to_dict() for e in user_expenses]
        meta, results = detector.fit_and_predict_batch(records)

        result_map = {r["id"]: r for r in results}
        for exp in user_expenses:
            res = result_map.get(exp.id)
            if res:
                exp.is_anomaly = res["is_anomaly"]
                exp.anomaly_score = res["anomaly_score"]
                exp.raw_score = res["raw_score"]
                exp.anomaly_label = res["anomaly_label"]
                exp.anomaly_reasons = json.dumps(res["explanation"])
                if exp.review_status not in ["reviewed", "dismissed"]:
                    exp.review_status = res["review_status"]

        db.session.commit()
        return meta

    # -------------------------------------------------------------
    # AUTHENTICATION ROUTES
    # -------------------------------------------------------------
    @app.route("/register", methods=["GET", "POST"])
    def register():
        if get_current_user():
            return redirect(url_for("index"))

        if request.method == "POST":
            username = request.form.get("username", "").strip()
            email = request.form.get("email", "").strip().lower()
            password = request.form.get("password", "")
            confirm = request.form.get("confirm_password", "")

            errors = []
            if not username or len(username) < 3 or len(username) > 30:
                errors.append("Username must be between 3 and 30 characters.")
            if not username.replace("_", "").isalnum():
                errors.append("Username can only contain letters, numbers, and underscores.")

            if not email or "@" not in email or "." not in email:
                errors.append("A valid email address is required.")

            if not password or len(password) < 6:
                errors.append("Password must be at least 6 characters long.")

            if password != confirm:
                errors.append("Passwords do not match.")

            # Uniqueness check
            if User.query.filter_by(username=username).first():
                errors.append("That username is already registered. Please choose another.")
            if User.query.filter_by(email=email).first():
                errors.append("That email address is already in use. Please sign in.")

            if errors:
                for err in errors:
                    flash(err, "danger")
                return render_template("auth/register.html", form_data=request.form)

            # Create user
            new_user = User(username=username, email=email)
            new_user.set_password(password)
            db.session.add(new_user)
            db.session.commit()

            # Set session
            session["user_id"] = new_user.id
            session["username"] = new_user.username
            flash(f"Welcome to SpendWise AI, {new_user.username}! Your account was created successfully.", "success")
            return redirect(url_for("index"))

        return render_template("auth/register.html", form_data={})

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if get_current_user():
            return redirect(url_for("index"))

        if request.method == "POST":
            identity = request.form.get("identity", "").strip()
            password = request.form.get("password", "")

            if not identity or not password:
                flash("Please provide both username/email and password.", "danger")
                return render_template("auth/login.html", form_data=request.form)

            user = (
                User.query.filter(
                    (User.username == identity) | (User.email == identity.lower())
                ).first()
            )

            if not user or not user.check_password(password):
                flash("Invalid credentials. Please verify your username/email and password.", "danger")
                return render_template("auth/login.html", form_data=request.form)

            # Valid login
            session["user_id"] = user.id
            session["username"] = user.username

            next_page = request.args.get("next")
            if next_page and next_page.startswith("/"):
                return redirect(next_page)
            return redirect(url_for("index"))

        return render_template("auth/login.html", form_data={})

    @app.route("/logout", methods=["POST"])
    def logout():
        session.clear()
        flash("You have been signed out safely.", "info")
        return redirect(url_for("login"))

    # -------------------------------------------------------------
    # 1. DASHBOARD (Protected & Scoped)
    # -------------------------------------------------------------
    @app.route("/")
    @login_required
    def index():
        current_u = get_current_user()
        year, month = get_selected_period(request)

        # Previous month calculation for comparison
        prev_month = 12 if month == 1 else month - 1
        prev_year = year - 1 if month == 1 else year

        # User-scoped month expenses
        month_expenses = (
            Expense.query.filter(
                Expense.user_id == current_u.id,
                extract("year", Expense.date) == year,
                extract("month", Expense.date) == month,
            )
            .order_by(Expense.date.desc())
            .all()
        )

        prev_month_expenses = (
            Expense.query.filter(
                Expense.user_id == current_u.id,
                extract("year", Expense.date) == prev_year,
                extract("month", Expense.date) == prev_month,
            )
            .all()
        )

        # Totals
        total_spent = sum(e.amount for e in month_expenses)
        prev_total_spent = sum(e.amount for e in prev_month_expenses)

        # User-scoped Budget
        budget_query = Budget.query.filter_by(year=year, month=month)
        budget_obj = budget_query.filter_by(user_id=current_u.id).first() or budget_query.filter_by(user_id=None).first()

        monthly_budget = budget_obj.amount if budget_obj else 0.0
        remaining_budget = monthly_budget - total_spent if monthly_budget > 0 else 0.0
        budget_used_percent = (
            round((total_spent / monthly_budget) * 100, 1) if monthly_budget > 0 else 0.0
        )

        # Anomalies
        anomaly_expenses = [e for e in month_expenses if e.is_anomaly]
        pending_alerts = [
            e for e in month_expenses if e.is_anomaly and e.review_status == "pending"
        ]

        # Month-over-Month change
        if prev_total_spent > 0:
            mom_change_percent = round(
                ((total_spent - prev_total_spent) / prev_total_spent) * 100, 1
            )
        else:
            mom_change_percent = 0.0

        # Category-wise breakdown
        category_totals = {cat: 0.0 for cat in app.config["EXPENSE_CATEGORIES"]}
        for e in month_expenses:
            category_totals[e.category] = category_totals.get(e.category, 0.0) + e.amount

        chart_categories = {k: v for k, v in category_totals.items() if v > 0}

        # Daily spending breakdown
        _, num_days = monthrange(year, month)
        daily_totals = {day: 0.0 for day in range(1, num_days + 1)}
        for e in month_expenses:
            daily_totals[e.date.day] = daily_totals.get(e.date.day, 0.0) + e.amount

        # Recent transactions (user-scoped)
        recent_expenses = (
            Expense.query.filter_by(user_id=current_u.id)
            .order_by(Expense.date.desc(), Expense.id.desc())
            .limit(6)
            .all()
        )

        return render_template(
            "index.html",
            selected_year=year,
            selected_month=month,
            total_spent=total_spent,
            prev_total_spent=prev_total_spent,
            mom_change_percent=mom_change_percent,
            monthly_budget=monthly_budget,
            remaining_budget=remaining_budget,
            budget_used_percent=budget_used_percent,
            anomalies_count=len(anomaly_expenses),
            pending_alerts=pending_alerts,
            recent_expenses=recent_expenses,
            chart_categories_json=json.dumps(chart_categories),
            daily_labels_json=json.dumps(list(daily_totals.keys())),
            daily_values_json=json.dumps(list(daily_totals.values())),
            month_expenses_count=len(month_expenses),
        )

    # -------------------------------------------------------------
    # 2. BUDGET MANAGEMENT (Protected & Scoped)
    # -------------------------------------------------------------
    @app.route("/budget/set", methods=["POST"])
    @login_required
    def set_budget():
        current_u = get_current_user()
        try:
            year = int(request.form.get("year", date.today().year))
            month = int(request.form.get("month", date.today().month))
            amount_val = float(request.form.get("amount", 0.0))

            if amount_val < 0:
                flash("Budget amount cannot be negative.", "danger")
                return redirect(url_for("index", year=year, month=month))

            # Query user-scoped budget or fallback
            budget = Budget.query.filter_by(year=year, month=month, user_id=current_u.id).first()
            if not budget:
                budget = Budget(year=year, month=month, user_id=current_u.id)
                db.session.add(budget)

            budget.amount = amount_val
            db.session.commit()
            flash(
                f"Budget of ₹{amount_val:,.2f} set successfully for {datetime(year, month, 1).strftime('%B %Y')}.",
                "success",
            )
        except ValueError:
            flash("Invalid budget amount or date provided.", "danger")

        return redirect(url_for("index", year=year, month=month))

    # -------------------------------------------------------------
    # 3. EXPENSE MANAGEMENT (CRUD - Protected & Scoped)
    # -------------------------------------------------------------
    @app.route("/expenses/add", methods=["GET", "POST"])
    @login_required
    def add_expense():
        current_u = get_current_user()
        if request.method == "POST":
            description = request.form.get("description", "").strip()
            amount_str = request.form.get("amount", "").strip()
            category = request.form.get("category", "").strip()
            date_str = request.form.get("date", "").strip()
            payment_method = request.form.get("payment_method", "").strip()
            notes = request.form.get("notes", "").strip()

            # Server-side Validation
            errors = []
            if not description or len(description) > 255:
                errors.append("Description is required and must be under 255 characters.")

            try:
                amount_val = float(amount_str)
                if amount_val <= 0:
                    errors.append("Amount must be greater than ₹0.")
                if amount_val > 10_000_000:
                    errors.append("Amount exceeds permitted maximum (₹1,00,00,000).")
            except (ValueError, TypeError):
                errors.append("A valid positive numerical amount is required.")

            if category not in app.config["EXPENSE_CATEGORIES"]:
                errors.append(f"Invalid category selected. Must be one of: {', '.join(app.config['EXPENSE_CATEGORIES'])}")

            if payment_method not in app.config["PAYMENT_METHODS"]:
                errors.append(f"Invalid payment method. Must be one of: {', '.join(app.config['PAYMENT_METHODS'])}")

            try:
                txn_date = datetime.strptime(date_str, "%Y-%m-%d").date()
            except (ValueError, TypeError):
                errors.append("Invalid date format. Expected YYYY-MM-DD.")

            if errors:
                for err in errors:
                    flash(err, "danger")
                return render_template(
                    "add_expense.html",
                    form_data=request.form,
                    today_str=date.today().strftime("%Y-%m-%d"),
                )

            # Create expense scoped to user
            new_expense = Expense(
                user_id=current_u.id,
                description=description,
                category=category,
                date=txn_date,
                payment_method=payment_method,
                notes=notes,
                is_demo=False,
            )
            new_expense.amount = amount_val
            db.session.add(new_expense)
            db.session.commit()

            # Evaluate with Anomaly Detector scoped to this user
            user_txns_count = Expense.query.filter_by(user_id=current_u.id).count()
            if user_txns_count >= app.config["MODEL_MIN_SAMPLES"]:
                retrain_model_on_database(current_u.id)
                db.session.refresh(new_expense)
                if new_expense.is_anomaly:
                    flash(
                        f"Expense of ₹{amount_val:,.2f} added. Notice: The AI model flagged this as potentially unusual spending for '{category}'.",
                        "warning",
                    )
                else:
                    flash(f"Expense of ₹{amount_val:,.2f} added successfully.", "success")
            else:
                needed = app.config["MODEL_MIN_SAMPLES"] - user_txns_count
                flash(
                    f"Expense of ₹{amount_val:,.2f} recorded. (Need {needed} more transaction(s) for personalized AI profiling).",
                    "success",
                )

            return redirect(url_for("history"))

        return render_template(
            "add_expense.html",
            form_data={},
            today_str=date.today().strftime("%Y-%m-%d"),
        )

    @app.route("/expenses/edit/<int:expense_id>", methods=["GET", "POST"])
    @login_required
    def edit_expense(expense_id):
        current_u = get_current_user()
        expense = Expense.query.get_or_404(expense_id)

        # Cross-user authorization check
        if expense.user_id is not None and expense.user_id != current_u.id:
            abort(403)

        if request.method == "POST":
            description = request.form.get("description", "").strip()
            amount_str = request.form.get("amount", "").strip()
            category = request.form.get("category", "").strip()
            date_str = request.form.get("date", "").strip()
            payment_method = request.form.get("payment_method", "").strip()
            notes = request.form.get("notes", "").strip()

            errors = []
            if not description or len(description) > 255:
                errors.append("Description is required and must be under 255 characters.")

            try:
                amount_val = float(amount_str)
                if amount_val <= 0:
                    errors.append("Amount must be greater than ₹0.")
            except (ValueError, TypeError):
                errors.append("A valid positive numerical amount is required.")

            if category not in app.config["EXPENSE_CATEGORIES"]:
                errors.append("Invalid category selected.")

            if payment_method not in app.config["PAYMENT_METHODS"]:
                errors.append("Invalid payment method.")

            try:
                txn_date = datetime.strptime(date_str, "%Y-%m-%d").date()
            except (ValueError, TypeError):
                errors.append("Invalid date format.")

            if errors:
                for err in errors:
                    flash(err, "danger")
                return render_template("edit_expense.html", expense=expense)

            expense.description = description
            expense.amount = amount_val
            expense.category = category
            expense.date = txn_date
            expense.payment_method = payment_method
            expense.notes = notes

            db.session.commit()

            # Refresh model predictions for this user
            if Expense.query.filter_by(user_id=current_u.id).count() >= app.config["MODEL_MIN_SAMPLES"]:
                retrain_model_on_database(current_u.id)

            flash(f"Expense #{expense.id} updated successfully.", "success")
            return redirect(url_for("history"))

        return render_template("edit_expense.html", expense=expense)

    @app.route("/expenses/delete/<int:expense_id>", methods=["POST"])
    @login_required
    def delete_expense(expense_id):
        current_u = get_current_user()
        expense = Expense.query.get_or_404(expense_id)

        # Cross-user authorization check
        if expense.user_id is not None and expense.user_id != current_u.id:
            abort(403)

        desc = expense.description
        amt = expense.amount
        db.session.delete(expense)
        db.session.commit()

        # Update model if enough samples
        if Expense.query.filter_by(user_id=current_u.id).count() >= app.config["MODEL_MIN_SAMPLES"]:
            retrain_model_on_database(current_u.id)

        flash(f"Expense '{desc}' (₹{amt:,.2f}) was successfully deleted.", "info")
        return redirect(request.referrer or url_for("history"))

    # -------------------------------------------------------------
    # 4. AI ANOMALY DETECTION (Protected & Scoped)
    # -------------------------------------------------------------
    @app.route("/anomalies")
    @login_required
    def anomalies():
        current_u = get_current_user()
        status_filter = request.args.get("status", "all")
        category_filter = request.args.get("category", "all")

        query = Expense.query.filter_by(user_id=current_u.id, is_anomaly=True)

        if status_filter != "all":
            query = query.filter_by(review_status=status_filter)
        if category_filter != "all":
            query = query.filter_by(category=category_filter)

        anomaly_records = query.order_by(Expense.date.desc(), Expense.anomaly_score.desc()).all()

        total_anomalies = Expense.query.filter_by(user_id=current_u.id, is_anomaly=True).count()
        pending_count = Expense.query.filter_by(user_id=current_u.id, is_anomaly=True, review_status="pending").count()
        reviewed_count = Expense.query.filter_by(user_id=current_u.id, is_anomaly=True, review_status="reviewed").count()
        dismissed_count = Expense.query.filter_by(user_id=current_u.id, is_anomaly=True, review_status="dismissed").count()

        user_txns_count = Expense.query.filter_by(user_id=current_u.id).count()

        return render_template(
            "anomalies.html",
            anomalies=anomaly_records,
            status_filter=status_filter,
            category_filter=category_filter,
            total_anomalies=total_anomalies,
            pending_count=pending_count,
            reviewed_count=reviewed_count,
            dismissed_count=dismissed_count,
            model_info={
                "is_fitted": detector.is_fitted and user_txns_count >= detector.min_samples,
                "samples_count": user_txns_count,
                "last_trained": detector.last_trained_at.strftime("%Y-%m-%d %H:%M:%S UTC")
                if detector.last_trained_at
                else "Not yet trained",
                "contamination": detector.contamination,
                "min_samples": detector.min_samples,
            },
        )

    @app.route("/anomalies/<int:expense_id>/review", methods=["POST"])
    @login_required
    def review_anomaly(expense_id):
        current_u = get_current_user()
        expense = Expense.query.get_or_404(expense_id)

        # Cross-user authorization check
        if expense.user_id is not None and expense.user_id != current_u.id:
            abort(403)

        new_status = request.form.get("status", "").strip().lower()
        if new_status in ["reviewed", "dismissed", "pending"]:
            expense.review_status = new_status
            db.session.commit()
            flash(
                f"Alert for '{expense.description}' marked as '{new_status.capitalize()}'.",
                "success",
            )
        else:
            flash("Invalid review status action.", "danger")

        return redirect(request.referrer or url_for("anomalies"))

    @app.route("/anomalies/retrain", methods=["POST"])
    @login_required
    def retrain_model():
        current_u = get_current_user()
        meta = retrain_model_on_database(current_u.id)
        if meta["status"] == "success":
            flash(
                f"IsolationForest model re-evaluated successfully across your {meta['sample_count']} transactions. Detected {meta['anomalies_detected']} unusual items.",
                "success",
            )
        elif meta["status"] == "insufficient_data":
            flash(meta["message"], "warning")
        else:
            flash("No data available to train anomaly model.", "info")

        return redirect(request.referrer or url_for("anomalies"))

    # -------------------------------------------------------------
    # 5. SPENDING ANALYTICS (Protected & Scoped)
    # -------------------------------------------------------------
    @app.route("/analytics")
    @login_required
    def analytics():
        current_u = get_current_user()
        category = request.args.get("category", "all")
        payment_method = request.args.get("payment_method", "all")
        start_date_str = request.args.get("start_date", "")
        end_date_str = request.args.get("end_date", "")

        query = Expense.query.filter_by(user_id=current_u.id)

        if category != "all":
            query = query.filter_by(category=category)
        if payment_method != "all":
            query = query.filter_by(payment_method=payment_method)

        if start_date_str:
            try:
                s_date = datetime.strptime(start_date_str, "%Y-%m-%d").date()
                query = query.filter(Expense.date >= s_date)
            except ValueError:
                pass

        if end_date_str:
            try:
                e_date = datetime.strptime(end_date_str, "%Y-%m-%d").date()
                query = query.filter(Expense.date <= e_date)
            except ValueError:
                pass

        filtered_expenses = query.order_by(Expense.date.asc()).all()

        total_amount = sum(e.amount for e in filtered_expenses)
        total_txns = len(filtered_expenses)
        avg_txn = total_amount / total_txns if total_txns > 0 else 0.0

        # Category spending
        cat_data = {cat: 0.0 for cat in app.config["EXPENSE_CATEGORIES"]}
        for e in filtered_expenses:
            cat_data[e.category] = cat_data.get(e.category, 0.0) + e.amount
        cat_data = {k: round(v, 2) for k, v in cat_data.items() if v > 0}

        # Payment method breakdown
        pm_data = {pm: 0.0 for pm in app.config["PAYMENT_METHODS"]}
        for e in filtered_expenses:
            pm_data[e.payment_method] = pm_data.get(e.payment_method, 0.0) + e.amount
        pm_data = {k: round(v, 2) for k, v in pm_data.items() if v > 0}

        # Monthly timeline aggregated
        month_map = {}
        for e in filtered_expenses:
            m_key = e.date.strftime("%Y-%m")
            month_map[m_key] = month_map.get(m_key, 0.0) + e.amount
        sorted_months = sorted(month_map.keys())

        # Daily timeline aggregated
        daily_map = {}
        for e in filtered_expenses:
            d_key = e.date.strftime("%Y-%m-%d")
            daily_map[d_key] = daily_map.get(d_key, 0.0) + e.amount
        sorted_days = sorted(daily_map.keys())

        return render_template(
            "analytics.html",
            total_amount=total_amount,
            total_txns=total_txns,
            avg_txn=avg_txn,
            category=category,
            payment_method=payment_method,
            start_date=start_date_str,
            end_date=end_date_str,
            has_data=total_txns > 0,
            cat_chart_json=json.dumps(cat_data),
            pm_chart_json=json.dumps(pm_data),
            monthly_labels_json=json.dumps(sorted_months),
            monthly_values_json=json.dumps([round(month_map[m], 2) for m in sorted_months]),
            daily_labels_json=json.dumps(sorted_days),
            daily_values_json=json.dumps([round(daily_map[d], 2) for d in sorted_days]),
        )

    # -------------------------------------------------------------
    # 6. TRANSACTION HISTORY (Protected & Scoped)
    # -------------------------------------------------------------
    @app.route("/history")
    @login_required
    def history():
        current_u = get_current_user()
        search_query = request.args.get("q", "").strip()
        category = request.args.get("category", "all")
        status = request.args.get("status", "all")
        sort_by = request.args.get("sort", "date_desc")
        page = request.args.get("page", 1, type=int)
        per_page = 15

        query = Expense.query.filter_by(user_id=current_u.id)

        if search_query:
            query = query.filter(
                (Expense.description.ilike(f"%{search_query}%"))
                | (Expense.notes.ilike(f"%{search_query}%"))
            )

        if category != "all":
            query = query.filter_by(category=category)

        if status == "anomalies":
            query = query.filter_by(is_anomaly=True)
        elif status == "normal":
            query = query.filter_by(is_anomaly=False)
        elif status in ["pending", "reviewed", "dismissed"]:
            query = query.filter_by(review_status=status)

        # Sorting
        if sort_by == "date_asc":
            query = query.order_by(Expense.date.asc(), Expense.id.asc())
        elif sort_by == "amount_desc":
            query = query.order_by(Expense.amount_paise.desc())
        elif sort_by == "amount_asc":
            query = query.order_by(Expense.amount_paise.asc())
        else:  # date_desc default
            query = query.order_by(Expense.date.desc(), Expense.id.desc())

        pagination = query.paginate(page=page, per_page=per_page, error_out=False)
        expenses = pagination.items

        return render_template(
            "history.html",
            expenses=expenses,
            pagination=pagination,
            search_query=search_query,
            selected_category=category,
            selected_status=status,
            selected_sort=sort_by,
        )

    # -------------------------------------------------------------
    # 7. MONTHLY REPORTS & CSV EXPORT (Protected & Scoped)
    # -------------------------------------------------------------
    @app.route("/reports")
    @login_required
    def reports():
        current_u = get_current_user()
        year, month = get_selected_period(request)

        # Previous month calculation
        prev_month = 12 if month == 1 else month - 1
        prev_year = year - 1 if month == 1 else year

        month_expenses = (
            Expense.query.filter(
                Expense.user_id == current_u.id,
                extract("year", Expense.date) == year,
                extract("month", Expense.date) == month,
            )
            .order_by(Expense.date.asc())
            .all()
        )

        prev_expenses = (
            Expense.query.filter(
                Expense.user_id == current_u.id,
                extract("year", Expense.date) == prev_year,
                extract("month", Expense.date) == prev_month,
            )
            .all()
        )

        total_spent = sum(e.amount for e in month_expenses)
        prev_total = sum(e.amount for e in prev_expenses)

        budget_query = Budget.query.filter_by(year=year, month=month)
        budget_obj = budget_query.filter_by(user_id=current_u.id).first() or budget_query.filter_by(user_id=None).first()

        monthly_budget = budget_obj.amount if budget_obj else 0.0
        budget_utilization = (
            round((total_spent / monthly_budget) * 100, 1) if monthly_budget > 0 else 0.0
        )

        highest_expense = max(month_expenses, key=lambda x: x.amount) if month_expenses else None
        flagged_count = sum(1 for e in month_expenses if e.is_anomaly)

        # Category breakdown
        category_breakdown = {}
        for cat in app.config["EXPENSE_CATEGORIES"]:
            cat_list = [e for e in month_expenses if e.category == cat]
            if cat_list:
                cat_sum = sum(e.amount for e in cat_list)
                pct = round((cat_sum / total_spent) * 100, 1) if total_spent > 0 else 0.0
                category_breakdown[cat] = {
                    "total": cat_sum,
                    "count": len(cat_list),
                    "percentage": pct,
                }

        month_name = datetime(year, month, 1).strftime("%B %Y")
        prev_month_name = datetime(prev_year, prev_month, 1).strftime("%B %Y")

        diff_amount = total_spent - prev_total
        diff_pct = round((diff_amount / prev_total) * 100, 1) if prev_total > 0 else 0.0

        return render_template(
            "reports.html",
            year=year,
            month=month,
            month_name=month_name,
            prev_month_name=prev_month_name,
            total_spent=total_spent,
            prev_total=prev_total,
            diff_amount=diff_amount,
            diff_pct=diff_pct,
            monthly_budget=monthly_budget,
            budget_utilization=budget_utilization,
            highest_expense=highest_expense,
            flagged_count=flagged_count,
            category_breakdown=category_breakdown,
            transactions_count=len(month_expenses),
        )

    @app.route("/reports/export-csv")
    @login_required
    def export_csv():
        """
        Exports user's monthly transactions into a clean CSV file.
        Secured against Formula Injection by prefixing risk tokens (=, +, -, @) with a single quote.
        """
        current_u = get_current_user()
        year, month = get_selected_period(request)

        expenses = (
            Expense.query.filter(
                Expense.user_id == current_u.id,
                extract("year", Expense.date) == year,
                extract("month", Expense.date) == month,
            )
            .order_by(Expense.date.asc())
            .all()
        )

        output = io.StringIO()
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)

        writer.writerow([
            "Transaction ID",
            "Date",
            "Description",
            "Category",
            "Amount (INR)",
            "Payment Method",
            "Anomaly Flag",
            "Unusualness Score (0-1)",
            "Review Status",
            "Notes",
            "Data Type",
        ])

        def sanitize_cell(val: Any) -> str:
            if val is None:
                return ""
            s = str(val).strip()
            if s and s[0] in ("=", "+", "-", "@", "\t", "\r"):
                return f"'{s}"
            return s

        for exp in expenses:
            writer.writerow([
                exp.id,
                format_csv_date(exp.date),
                sanitize_cell(exp.description),
                sanitize_cell(exp.category),
                f"{exp.amount:.2f}",
                sanitize_cell(exp.payment_method),
                "Potentially Unusual" if exp.is_anomaly else "Typical",
                f"{exp.anomaly_score:.3f}" if exp.anomaly_score is not None else "N/A",
                sanitize_cell(exp.review_status),
                sanitize_cell(exp.notes),
                "Demonstration" if exp.is_demo else "User Entered",
            ])

        output.seek(0)
        filename = f"SpendWise_Report_{year}_{month:02d}.csv"
        return Response(
            output.getvalue(),
            mimetype="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )

    @app.route("/reports/export-excel")
    @login_required
    def export_excel():
        """
        Exports user's monthly transactions into a formatted Microsoft Excel (.xlsx) workbook.
        Provides explicit YYYY-MM-DD date cells, generous column widths (preventing Excel
        '#####' display), numeric currency formatting, and formula-injection defenses.
        """
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        current_u = get_current_user()
        year, month = get_selected_period(request)

        expenses = (
            Expense.query.filter(
                Expense.user_id == current_u.id,
                extract("year", Expense.date) == year,
                extract("month", Expense.date) == month,
            )
            .order_by(Expense.date.asc())
            .all()
        )

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{year}-{month:02d} Expenses"
        ws.views.sheetView[0].showGridLines = True

        headers = [
            "Transaction ID",
            "Date",
            "Description",
            "Category",
            "Amount (INR)",
            "Payment Method",
            "Anomaly Flag",
            "Unusualness Score (0-1)",
            "Review Status",
            "Notes",
            "Data Type",
        ]

        def sanitize_cell(val: Any) -> str:
            if val is None:
                return ""
            s = str(val).strip()
            if s and s[0] in ("=", "+", "-", "@", "\t", "\r"):
                return f"'{s}"
            return s

        # Header styling
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="312E81", end_color="312E81", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # Data styling
        data_font = Font(name="Segoe UI", size=9)
        thin_border = Border(
            left=Side(style="thin", color="E2E8F0"),
            right=Side(style="thin", color="E2E8F0"),
            top=Side(style="thin", color="E2E8F0"),
            bottom=Side(style="thin", color="E2E8F0"),
        )

        for row_idx, exp in enumerate(expenses, start=2):
            date_val = exp.date if isinstance(exp.date, date) else datetime.strptime(format_csv_date(exp.date), "%Y-%m-%d").date()

            row_data = [
                exp.id,
                date_val,
                sanitize_cell(exp.description),
                sanitize_cell(exp.category),
                round(exp.amount, 2),
                sanitize_cell(exp.payment_method),
                "Potentially Unusual" if exp.is_anomaly else "Typical",
                round(exp.anomaly_score, 3) if exp.anomaly_score is not None else "N/A",
                sanitize_cell(exp.review_status),
                sanitize_cell(exp.notes),
                "Demonstration" if exp.is_demo else "User Entered",
            ]
            ws.append(row_data)

            # Alignment & number formats
            ws.cell(row=row_idx, column=1).alignment = Alignment(horizontal="center")
            date_cell = ws.cell(row=row_idx, column=2)
            date_cell.number_format = "yyyy-mm-dd"
            date_cell.alignment = Alignment(horizontal="center")
            amount_cell = ws.cell(row=row_idx, column=5)
            amount_cell.number_format = "#,##0.00"
            amount_cell.alignment = Alignment(horizontal="right")
            score_cell = ws.cell(row=row_idx, column=8)
            if isinstance(score_cell.value, float):
                score_cell.number_format = "0.000"
                score_cell.alignment = Alignment(horizontal="right")
            else:
                score_cell.alignment = Alignment(horizontal="center")

            for col_num in range(1, len(headers) + 1):
                c = ws.cell(row=row_idx, column=col_num)
                c.font = data_font
                c.border = thin_border

        # Ample column widths to eliminate Excel '#####' display completely
        column_widths = {
            1: 16,
            2: 15,
            3: 32,
            4: 18,
            5: 18,
            6: 20,
            7: 22,
            8: 24,
            9: 16,
            10: 36,
            11: 18,
        }
        for col_idx, width in column_widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"SpendWise_Report_{year}_{month:02d}.xlsx"
        return Response(
            output.getvalue(),
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )

    # -------------------------------------------------------------
    # 8. DEMONSTRATION DATA (Protected & User-Scoped)
    # -------------------------------------------------------------
    @app.route("/demo/seed", methods=["POST"])
    @login_required
    def seed_demo():
        current_u = get_current_user()
        from data.generate_demo_data import seed_demo_data
        result = seed_demo_data(user_id=current_u.id)
        flash(
            f"Demonstration data loaded! Added {result['seeded_count']} realistic transactions with {result['anomalies_flagged']} flagged anomaly scenarios.",
            "success",
        )
        return redirect(request.referrer or url_for("index"))

    @app.route("/demo/clear", methods=["POST"])
    @login_required
    def clear_demo():
        current_u = get_current_user()
        from data.generate_demo_data import clear_demo_data
        deleted = clear_demo_data(user_id=current_u.id)
        if Expense.query.filter_by(user_id=current_u.id).count() >= app.config["MODEL_MIN_SAMPLES"]:
            retrain_model_on_database(current_u.id)
        flash(f"Removed {deleted} demonstration transactions from your account.", "info")
        return redirect(request.referrer or url_for("index"))

    return app


# Default application instance
app = create_app()

if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=False)
