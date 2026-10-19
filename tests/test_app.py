"""
tests/test_app.py - Integration and Route Tests for SpendWise AI
================================================================
Uses an isolated in-memory SQLite database (sqlite:///:memory:) via TestingConfig.
Tests startup, CRUD workflows, server-side validation, budget mechanics,
review status transitions, formula-injection sanitization, and error pages.
"""

import csv
from datetime import date, datetime
import io
import re
import pytest
from app import create_app, format_csv_date
from config import TestingConfig
from database import db, Expense, Budget, User


@pytest.fixture
def app():
    """Create a test application instance with in-memory database."""
    test_app = create_app(TestingConfig)
    with test_app.app_context():
        db.create_all()
        yield test_app
        db.session.remove()
        db.drop_all()


@pytest.fixture
def auth_user(app):
    """Seed a default test user."""
    with app.app_context():
        user = User(username="test_tester", email="test@example.com")
        user.set_password("securepassword123")
        db.session.add(user)
        db.session.commit()
        return user.id


@pytest.fixture
def client(app, auth_user):
    """Test HTTP client with active session."""
    client = app.test_client()
    with client.session_transaction() as sess:
        sess["user_id"] = auth_user
        sess["username"] = "test_tester"
    return client


def test_app_startup_and_empty_dashboard(client):
    """
    Startup & Empty DB Test:
    Verify dashboard loads HTTP 200 with zero records and renders empty state cleanly.
    """
    response = client.get("/")
    assert response.status_code == 200
    assert b"SpendWise AI" in response.data
    assert b"Financial Dashboard" in response.data
    assert b"No Transactions Recorded Yet" in response.data


def test_add_expense_valid(client, app):
    """
    CRUD Add Test:
    Add a valid expense through POST and verify it is persisted in the database.
    """
    payload = {
        "description": "Textbooks for Semester",
        "amount": "1450.50",
        "category": "Education",
        "date": "2026-10-15",
        "payment_method": "UPI",
        "notes": "College book store",
    }
    response = client.post("/expenses/add", data=payload, follow_redirects=True)
    assert response.status_code == 200

    with app.app_context():
        exp = Expense.query.first()
        assert exp is not None
        assert exp.description == "Textbooks for Semester"
        assert exp.amount == 1450.50
        assert exp.amount_paise == 145050
        assert exp.category == "Education"
        assert exp.payment_method == "UPI"


def test_add_expense_invalid_validation(client, app):
    """
    Input Validation Test:
    Server must reject negative amounts, empty descriptions, or invalid categories.
    """
    # Negative amount
    bad_payload = {
        "description": "Invalid Negative",
        "amount": "-500",
        "category": "Food",
        "date": "2026-10-15",
        "payment_method": "Cash",
    }
    res = client.post("/expenses/add", data=bad_payload)
    assert res.status_code == 200
    assert b"Amount must be greater than" in res.data

    # Invalid Category
    bad_cat_payload = {
        "description": "Invalid Category",
        "amount": "500",
        "category": "FakeCategory123",
        "date": "2026-10-15",
        "payment_method": "Cash",
    }
    res2 = client.post("/expenses/add", data=bad_cat_payload)
    assert res2.status_code == 200
    assert b"Invalid category selected" in res2.data

    with app.app_context():
        assert Expense.query.count() == 0


def test_edit_expense(client, app, auth_user):
    """
    CRUD Edit Test:
    Create an expense, update it via POST, and verify fields change properly.
    """
    with app.app_context():
        exp = Expense(
            user_id=auth_user,
            description="Grocery Shopping",
            category="Food",
            date=date(2026, 10, 10),
            payment_method="Cash",
        )
        exp.amount = 800.0
        db.session.add(exp)
        db.session.commit()
        exp_id = exp.id

    edit_payload = {
        "description": "Supermarket Monthly Grocery",
        "amount": "1200.00",
        "category": "Shopping",
        "date": "2026-10-11",
        "payment_method": "Debit Card",
        "notes": "Updated note",
    }
    res = client.post(f"/expenses/edit/{exp_id}", data=edit_payload, follow_redirects=True)
    assert res.status_code == 200

    with app.app_context():
        updated = db.session.get(Expense, exp_id)
        assert updated.description == "Supermarket Monthly Grocery"
        assert updated.amount == 1200.00
        assert updated.category == "Shopping"
        assert updated.payment_method == "Debit Card"


def test_delete_expense(client, app, auth_user):
    """
    CRUD Delete Test:
    Delete an existing expense and verify removal from database.
    """
    with app.app_context():
        exp = Expense(
            user_id=auth_user,
            description="Coffee",
            category="Food",
            date=date(2026, 10, 12),
            payment_method="UPI",
        )
        exp.amount = 120.0
        db.session.add(exp)
        db.session.commit()
        exp_id = exp.id

    res = client.post(f"/expenses/delete/{exp_id}", follow_redirects=True)
    assert res.status_code == 200

    with app.app_context():
        assert db.session.get(Expense, exp_id) is None


def test_budget_setting_and_calculations(client, app, auth_user):
    """
    Budget Mechanics Test:
    Set monthly budget, add expense in that month, and check remaining balance.
    """
    payload = {"year": "2026", "month": "10", "amount": "30000.00"}
    res = client.post("/budget/set", data=payload, follow_redirects=True)
    assert res.status_code == 200
    assert b"Budget of" in res.data

    with app.app_context():
        b = Budget.query.filter_by(year=2026, month=10).first()
        assert b is not None
        assert b.amount == 30000.00

        # Add expense of 5,000
        e = Expense(
            user_id=auth_user,
            description="Rent",
            category="Bills",
            date=date(2026, 10, 5),
            payment_method="Net Banking",
        )
        e.amount = 5000.00
        db.session.add(e)
        db.session.commit()

    # Check dashboard shows budget and remaining calculation
    dash = client.get("/?year=2026&month=10")
    assert dash.status_code == 200
    assert b"30000.00" in dash.data
    assert b"25000.00" in dash.data


def test_anomaly_review_workflow(client, app, auth_user):
    """
    Anomaly Review Action Test:
    Flagged anomaly can be marked as 'reviewed' or 'dismissed'.
    """
    with app.app_context():
        exp = Expense(
            user_id=auth_user,
            description="Suspicious Charge",
            category="Shopping",
            date=date(2026, 10, 18),
            payment_method="Credit Card",
            is_anomaly=True,
            review_status="pending",
        )
        exp.amount = 55000.0
        db.session.add(exp)
        db.session.commit()
        exp_id = exp.id

    # Mark as reviewed
    res = client.post(
        f"/anomalies/{exp_id}/review",
        data={"status": "reviewed"},
        follow_redirects=True,
    )
    assert res.status_code == 200

    with app.app_context():
        reviewed_exp = db.session.get(Expense, exp_id)
        assert reviewed_exp.review_status == "reviewed"


def test_csv_export_and_formula_injection_defense(client, app, auth_user):
    """
    Security Test - CSV Formula Injection Mitigation:
    A description or note beginning with dangerous formula triggers (=, +, -, @)
    must be escaped with a single quote (') to prevent Excel macro execution.
    """
    with app.app_context():
        dangerous_exp = Expense(
            user_id=auth_user,
            description="=cmd|' /C calc'!A0",
            category="Food",
            date=date(2026, 10, 14),
            payment_method="Cash",
            notes="+SUM(1+1)*cmd",
        )
        dangerous_exp.amount = 450.0
        db.session.add(dangerous_exp)
        db.session.commit()

    res = client.get("/reports/export-csv?year=2026&month=10")
    assert res.status_code == 200
    assert res.mimetype == "text/csv"

    csv_content = res.data.decode("utf-8")
    # Must be escaped with leading apostrophe
    assert "'=cmd|" in csv_content
    assert "'+SUM(1+1)" in csv_content


def test_csv_export_date_format_and_column_integrity(client, app, auth_user):
    """
    CSV Export Date Formatting & Column Integrity Test:
    1. Verifies format_csv_date helper handles date objects, datetimes, ISO strings,
       and Excel date serial numbers (such as 46926) converting to YYYY-MM-DD.
    2. Verifies exported CSV contains human-readable YYYY-MM-DD date strings rather than
       raw numbers or serial numbers.
    3. Verifies all 11 columns, headers, amounts, categories, and methods are preserved.
    4. Verifies multi-user data isolation (another user cannot export user's records).
    """
    # 1. Helper unit assertions
    assert format_csv_date(date(2026, 10, 19)) == "2026-10-19"
    assert format_csv_date(datetime(2026, 10, 19, 14, 30)) == "2026-10-19"
    assert format_csv_date("2026-10-19") == "2026-10-19"
    assert format_csv_date(46926) == "2028-06-22"
    assert format_csv_date("46926") == "2028-06-22"
    assert format_csv_date(46314) == "2026-10-19"
    assert format_csv_date(None) == ""

    # 2. Add sample expense for auth_user
    with app.app_context():
        exp = Expense(
            user_id=auth_user,
            description="Organic Groceries",
            category="Food",
            date=date(2026, 10, 19),
            payment_method="UPI",
            notes="Weekly grocery shopping",
        )
        exp.amount = 1250.50
        db.session.add(exp)
        db.session.commit()

    # 3. Export CSV
    res = client.get("/reports/export-csv?year=2026&month=10")
    assert res.status_code == 200
    assert res.mimetype == "text/csv"

    csv_text = res.data.decode("utf-8")
    reader = list(csv.reader(io.StringIO(csv_text)))

    # Assert Headers
    expected_headers = [
        "Transaction ID",
        "Date",
        "Description",
        "Category",
        "Amount (INR)",
        "Payment Method",
        "Anomaly Flag",
        "Unusualness Score (0-1)",
        "Review Status",
        "Notes",
        "Data Type",
    ]
    assert reader[0] == expected_headers

    # Find the row for "Organic Groceries"
    row = next((r for r in reader[1:] if "Organic Groceries" in r[2]), None)
    assert row is not None, "Expense record should be present in CSV export"

    # Check Date format specifically
    date_val = row[1]
    assert date_val == "2026-10-19"
    assert re.match(r"^\d{4}-\d{2}-\d{2}$", date_val), f"Expected YYYY-MM-DD, got {date_val}"
    # Verify it is not an Excel date serial number or raw number
    assert not date_val.isdigit()
    assert date_val != "46314"
    assert date_val != "46926"

    # Check remaining fields
    assert row[2] == "Organic Groceries"
    assert row[3] == "Food"
    assert row[4] == "1250.50"
    assert row[5] == "UPI"
    assert row[6] in ("Typical", "Potentially Unusual")
    assert row[9] == "Weekly grocery shopping"
    assert row[10] == "User Entered"

    # Add a demo expense to test "Demonstration" label in CSV
    with app.app_context():
        demo_exp = Expense(
            user_id=auth_user,
            description="Synthetic Dining Demo",
            category="Food",
            date=date(2026, 10, 20),
            payment_method="Credit Card",
            is_demo=True,
        )
        demo_exp.amount = 3500.0
        db.session.add(demo_exp)
        db.session.commit()
        demo_exp_id = demo_exp.id

    res_demo = client.get("/reports/export-csv?year=2026&month=10")
    demo_reader = list(csv.reader(io.StringIO(res_demo.data.decode("utf-8"))))
    demo_row = next((r for r in demo_reader[1:] if "Synthetic Dining Demo" in r[2]), None)
    assert demo_row is not None
    assert demo_row[10] == "Demonstration"

    # Delete the demo expense and verify it is NO LONGER exported in CSV
    with app.app_context():
        d_exp = db.session.get(Expense, demo_exp_id)
        db.session.delete(d_exp)
        db.session.commit()

    res_after_del = client.get("/reports/export-csv?year=2026&month=10")
    assert "Synthetic Dining Demo" not in res_after_del.data.decode("utf-8")

    # 4. Multi-tenant isolation test: User B cannot export auth_user's CSV records
    with app.app_context():
        other_user = User(username="other_csv_user", email="other_csv@test.com")
        other_user.set_password("SecurePass99")
        db.session.add(other_user)
        db.session.commit()

    client_b = app.test_client()
    client_b.post("/login", data={"identity": "other_csv_user", "password": "SecurePass99"})
    res_b = client_b.get("/reports/export-csv?year=2026&month=10")
    assert res_b.status_code == 200
    reader_b = list(csv.reader(io.StringIO(res_b.data.decode("utf-8"))))
    # Reader B only has the header row (no records from auth_user)
    assert len(reader_b) == 1
    assert "Organic Groceries" not in res_b.data.decode("utf-8")


def test_excel_export_and_formatting(client, app, auth_user):
    """
    Excel Export (.xlsx) Test:
    1. Tests /reports/export-excel returns valid application/vnd.openxmlformats spreadsheet.
    2. Validates all 11 columns match exact headers.
    3. Confirms Date cell is formatted with number_format 'yyyy-mm-dd' and column width >= 14.
    4. Confirms Amount cell is formatted as '#,##0.00'.
    5. Confirms formula injection prevention in cells.
    6. Verifies cross-user isolation.
    """
    import openpyxl

    with app.app_context():
        exp = Expense(
            user_id=auth_user,
            description="=dangerous_calc()",
            category="Food",
            date=date(2026, 10, 25),
            payment_method="Cash",
            notes="+cmd|' /C test'!A0",
        )
        exp.amount = 850.75
        db.session.add(exp)
        db.session.commit()

    res = client.get("/reports/export-excel?year=2026&month=10")
    assert res.status_code == 200
    assert res.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "SpendWise_Report_2026_10.xlsx" in res.headers["Content-Disposition"]

    # Open workbook in-memory
    wb = openpyxl.load_workbook(io.BytesIO(res.data))
    ws = wb.active
    assert ws.title == "2026-10 Expenses"

    # Verify headers
    expected_headers = [
        "Transaction ID",
        "Date",
        "Description",
        "Category",
        "Amount (INR)",
        "Payment Method",
        "Anomaly Flag",
        "Unusualness Score (0-1)",
        "Review Status",
        "Notes",
        "Data Type",
    ]
    header_values = [ws.cell(row=1, column=col).value for col in range(1, 12)]
    assert header_values == expected_headers

    # Verify data row exists
    assert ws.max_row >= 2
    row_date_cell = ws.cell(row=2, column=2)
    assert row_date_cell.number_format == "yyyy-mm-dd"
    # Column width for Date must be >= 14 to prevent #####
    assert ws.column_dimensions["B"].width >= 14

    # Amount format
    row_amount_cell = ws.cell(row=2, column=5)
    assert row_amount_cell.number_format == "#,##0.00"

    # Formula injection escaping
    row_desc = ws.cell(row=2, column=3).value
    assert str(row_desc).startswith("'=")
    row_notes = ws.cell(row=2, column=10).value
    assert str(row_notes).startswith("'+")

    # Multi-tenant isolation for Excel
    with app.app_context():
        other_user = User(username="excel_other_user", email="excel_other@test.com")
        other_user.set_password("ExcelPass123")
        db.session.add(other_user)
        db.session.commit()

    client_other = app.test_client()
    client_other.post("/login", data={"identity": "excel_other_user", "password": "ExcelPass123"})
    res_other = client_other.get("/reports/export-excel?year=2026&month=10")
    wb_other = openpyxl.load_workbook(io.BytesIO(res_other.data))
    ws_other = wb_other.active
    # Only header row should exist
    assert ws_other.max_row == 1


def test_safe_404_handling(client):
    """
    Security & Routing Test:
    Non-existent route should return customized 404 template without stack trace.
    """
    res = client.get("/non-existent-endpoint-abc")
    assert res.status_code == 404
    assert b"404" in res.data
    assert b"Page Not Found" in res.data


def test_demo_data_seeding_and_clearing(client, app, auth_user):
    """
    Regression Test:
    Verify that POST /demo/seed executes successfully without 500 Internal Server Error,
    scopes all seeded transactions to the logged-in user, and POST /demo/clear removes them.
    """
    # 1. Seed Demo Data
    res = client.post("/demo/seed", follow_redirects=True)
    assert res.status_code == 200
    assert b"Demonstration data loaded!" in res.data

    with app.app_context():
        # Confirm expenses exist and are linked to auth_user
        demo_expenses = Expense.query.filter_by(user_id=auth_user, is_demo=True).all()
        assert len(demo_expenses) > 30, "Expected at least 30 demo expenses to be generated"
        # Confirm anomaly flags are populated
        flagged = [e for e in demo_expenses if e.is_anomaly]
        assert len(flagged) > 0, "Expected demo anomalies to be flagged"

    # 2. Clear Demo Data
    clear_res = client.post("/demo/clear", follow_redirects=True)
    assert clear_res.status_code == 200
    assert b"Removed" in clear_res.data

    with app.app_context():
        remaining = Expense.query.filter_by(user_id=auth_user, is_demo=True).count()
        assert remaining == 0


def test_monthly_reports_export_buttons_ui(client):
    """
    UI Test:
    Verify that the Monthly Reports page renders the Excel export button
    and that the CSV export button is removed.
    """
    res = client.get("/reports?year=2026&month=10")
    assert res.status_code == 200
    # Excel export button is present
    assert b"Export Excel (.xlsx)" in res.data
    assert b"/reports/export-excel?year=2026&amp;month=10" in res.data or b"/reports/export-excel" in res.data
    # CSV export button is removed
    assert b"/reports/export-csv" not in res.data


